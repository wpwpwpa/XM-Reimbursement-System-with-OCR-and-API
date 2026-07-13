"""报销整理工作流 —— 机械后端（AI 视觉读取后的落盘脚本）。

输入：AI 按 SKILL.md 第 7 节产出的 records.json
输出：
  1) 报销导入_{费用类型}_{时间戳}.xlsx  （套模板，Row1~Row4 固定结构）
  2) 订单汇总_{时间戳}.xlsx
  3) 按规则重命名源文件（规则 A 默认 / 规则 B 带时间）

仅依赖 openpyxl。不引用 XM报销OCR 项目任何代码，可独立拷走使用。

用法：
  pip install openpyxl
  python build_outputs.py --records records.json --out-dir OUT [--taobao-excel 淘宝.xlsx]
  # 可选 --taobao-excel：对缺 order_time/product_name 的成功截图，按价格自动匹配补全
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# 金额 / 文件名格式化
# ---------------------------------------------------------------------------
def fmt_amount(a) -> str:
    """整数去尾部 .0（14 -> '14'），小数保留实际位数（26.52 / 26.5）。"""
    a = float(a)
    if a == int(a):
        return str(int(a))
    return f"{a:.2f}".rstrip("0").rstrip(".")


def new_filename(rec: dict, rule: str) -> str | None:
    """按规则生成新文件名（不含路径）。失败记录返回 None（不改名）。"""
    if rec.get("status") != "success":
        return None
    src = Path(rec["original_path"])
    ext = src.suffix  # 含点，如 .png

    if rec.get("source") == "invoice":
        date = rec.get("order_time") or ""
        amt = rec.get("amount")
        seller = rec.get("seller") or "未知销售方"
        if amt is None:
            return None
        return f"{date}-{fmt_amount(amt)}元-{seller}-发票{ext}"

    amount = rec.get("amount")
    platform = rec.get("platform")
    if amount is None or not platform:
        return None  # 金额或平台缺失 → 不改名

    if rule == "B":
        t = rec.get("order_time") or ""
        return f"{t}-{fmt_amount(amount)}-{platform}{ext}"
    # 规则 A（默认）
    return f"{fmt_amount(amount)}-{platform}{ext}"


# ---------------------------------------------------------------------------
# 淘宝 Excel 读取 + 按价格匹配（可选补全）
# ---------------------------------------------------------------------------
COLUMN_ALIASES = {
    "price": ["实付金额", "价格", "金额", "实付款", "付款金额", "成交金额",
              "订单金额", "实付"],
    "order_time": ["下单时间", "下单日期", "付款时间", "成交时间", "订单时间",
                   "创建时间", "购买时间", "成交日期", "订单提交时间",
                   "提交时间", "订单创建时间", "拍下时间"],
    "product": ["商品名称", "宝贝标题", "商品", "标题", "商品标题", "商品信息"],
}
PRICE_TOLERANCE = 5.0


def _match_header(headers, keywords):
    best_idx, best_len = -1, 0
    for idx, h in enumerate(headers):
        hs = str(h).strip()
        for kw in keywords:
            if kw in hs and len(kw) > best_len:
                best_idx, best_len = idx, len(kw)
    return best_idx


def _to_float(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("¥", "").replace("￥", "").replace(",", "").strip()
    m = re.search(r"\d+(?:\.\d+)?", s)
    return float(m.group()) if m else None


def _norm_date(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    m = re.search(r"(\d{4})[-/年.](\d{1,2})[-/月.](\d{1,2})", s)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    return s


def load_taobao_excel(path: str) -> dict:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        try:
            headers = list(next(it))
        except StopIteration:
            return {"rows": [], "headers": []}
        idx = {k: _match_header(headers, kws) for k, kws in COLUMN_ALIASES.items()}
        rows = []
        for r in it:
            if r is None or all(c is None for c in r):
                continue
            rows.append({
                "price": _to_float(r[idx["price"]]) if idx["price"] >= 0 else None,
                "order_time": _norm_date(r[idx["order_time"]]) if idx["order_time"] >= 0 else None,
                "product": str(r[idx["product"]]).strip()
                if (idx["product"] >= 0 and r[idx["product"]] is not None) else None,
            })
        return {"rows": rows, "headers": [str(h) for h in headers]}
    finally:
        wb.close()


def match_by_price(excel_data: dict, price: float | None):
    """返回 (row_or_None, ambiguous)。价格容差 ±5 元取最近一行。"""
    if price is None:
        return None, False
    cands = [(abs(r["price"] - price), r) for r in excel_data["rows"]
             if r["price"] is not None and abs(r["price"] - price) <= PRICE_TOLERANCE]
    if not cands:
        return None, False
    cands.sort(key=lambda x: x[0])
    ambiguous = len(cands) > 1 and (cands[1][0] - cands[0][0]) < PRICE_TOLERANCE * 0.5
    return cands[0][1], ambiguous


def fill_from_taobao(records: list[dict], excel_path: str) -> list[str]:
    """对缺 order_time/product_name 的成功截图，按价格匹配补全。返回歧义提示列表。"""
    data = load_taobao_excel(excel_path)
    warnings = []
    for rec in records:
        if rec.get("source") != "screenshot" or rec.get("status") != "success":
            continue
        if rec.get("order_time") and rec.get("product_name"):
            continue
        row, amb = match_by_price(data, rec.get("amount"))
        if row:
            if not rec.get("order_time"):
                rec["order_time"] = row.get("order_time") or ""
            if not rec.get("product_name"):
                rec["product_name"] = row.get("product") or ""
            if amb:
                warnings.append(
                    f"价格 {rec.get('amount')} 匹配到淘宝多行，已取最近一行，请核对："
                    f"{Path(rec['original_path']).name}")
    return warnings


# ---------------------------------------------------------------------------
# 报销导入 Excel
# ---------------------------------------------------------------------------
def _build_reason(order_time, platform, product, amount) -> str:
    md = ""
    if order_time:
        try:
            _, mo, d = order_time.split("-")
            md = f"{int(mo)}月{int(d)}日"
        except Exception:
            md = ""
    date_part = f"{md}，" if md else ""
    buy_part = f"{platform}购买" if platform else "购买"
    if product:
        buy_part = f"{buy_part}{product}等"
    amt_str = f"{fmt_amount(amount)}元" if amount is not None else "元"
    return f"{date_part}{buy_part}，共计{amt_str}。"


def _clean_header(h) -> str:
    return str(h).split("[")[0].strip() if h is not None else ""


def write_reimbursement_excel(template_path: str | None, expense_type: str,
                              rows: list[dict], out_path: str) -> str:
    new_wb = Workbook()
    ws = new_wb.active
    ws.title = expense_type

    instruction = None
    sys_id = None
    headers = ["金额", "消费日期", "发票形式", "发票", "消费事由"]

    if template_path and Path(template_path).exists():
        src = load_workbook(template_path).active
        instruction = src["A1"].value
        sys_id = src["A3"].value
        hs = [_clean_header(src.cell(row=4, column=c).value)
              for c in range(1, src.max_column + 1)]
        hs = [h for h in hs if h]
        if hs:
            headers = hs

    if instruction is not None:
        ws["A1"] = instruction
    ws["A2"] = expense_type
    if sys_id is not None:
        ws["A3"] = sys_id
    for c, h in enumerate(headers, start=1):
        ws.cell(row=4, column=c, value=h)

    for r_idx, row in enumerate(rows, start=5):
        for c, h in enumerate(headers, start=1):
            ws.cell(row=r_idx, column=c, value=row.get(h, ""))

    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 22

    out_path = str(Path(out_path).with_suffix(".xlsx"))
    new_wb.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# 订单汇总 Excel
# ---------------------------------------------------------------------------
def write_order_excel(rows: list[dict], out_path: str) -> str:
    headers = ["时间", "商品名称", "价格", "平台", "原文件名"]
    wb = Workbook()
    ws = wb.active
    ws.title = "订单汇总"
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
    widths = [max(12, min(40, len(h) * 3)) for h in headers]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    out_path = str(Path(out_path).with_suffix(".xlsx"))
    wb.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--records", required=True, help="records.json 路径")
    ap.add_argument("--out-dir", required=True, help="输出文件夹")
    ap.add_argument("--taobao-excel", default=None, help="可选：淘宝订单Excel，按价格补全字段")
    args = ap.parse_args()

    rec_path = Path(args.records)
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    records = json.loads(rec_path.read_text(encoding="utf-8"))

    rule = records.get("rename_rule", "A")
    expense_type = records.get("expense_type", "买耗材")
    template_path = records.get("template_path")

    # 可选：淘宝 Excel 补全
    if args.taobao_excel:
        for w in fill_from_taobao(records["records"], args.taobao_excel):
            print("⚠️", w)

    # 1) 重命名源文件
    renamed, skipped, conflicts = [], [], []
    for rec in records["records"]:
        new_name = new_filename(rec, rule)
        if not new_name:
            if rec.get("status") != "success":
                skipped.append(Path(rec["original_path"]).name)
            continue
        src = Path(rec["original_path"])
        dst = src.parent / new_name
        k = 1
        while dst.exists():
            stem = src.parent / (new_name[:new_name.rfind('.')] + f"({k})" + src.suffix)
            dst = stem
            k += 1
            if k > 2:
                conflicts.append(new_name)
        try:
            shutil.move(str(src), str(dst))
            renamed.append((src.name, dst.name))
        except Exception as e:
            skipped.append(f"{src.name}（重命名失败：{e}）")

    # 2) 报销导入 Excel（仅 success 且含 金额+消费日期）
    reimb_rows = []
    for rec in records["records"]:
        if rec.get("status") != "success":
            continue
        amount = rec.get("amount")
        order_time = rec.get("order_time")
        if amount is None or not order_time:
            continue
        reimb_rows.append({
            "金额": amount,
            "消费日期": order_time,
            "发票形式": "",
            "发票": 1 if rec.get("source") == "invoice" else 0,
            "消费事由": _build_reason(order_time, rec.get("platform") or "",
                                     rec.get("product_name"), amount),
        })
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    reimb_path = write_reimbursement_excel(
        template_path, expense_type, reimb_rows,
        str(out_dir / f"报销导入_{expense_type}_{ts}"))

    # 3) 订单汇总 Excel（全部记录）
    order_rows = [{
        "时间": rec.get("order_time") or "",
        "商品名称": rec.get("product_name") or "",
        "价格": rec.get("amount") if rec.get("amount") is not None else "",
        "平台": rec.get("platform") or "",
        "原文件名": Path(rec["original_path"]).name,
    } for rec in records["records"]]
    order_path = write_order_excel(order_rows, str(out_dir / f"订单汇总_{ts}"))

    # 汇总
    succ = sum(1 for r in records["records"] if r.get("status") == "success")
    fail = sum(1 for r in records["records"] if r.get("status") != "success")
    print(f"✅ 重命名成功：{len(renamed)} 个")
    for o, n in renamed:
        print(f"   {o}  →  {n}")
    if skipped:
        print(f"⏭️ 跳过（失败/未改名）：{len(skipped)} 个")
        for s in skipped:
            print(f"   - {s}")
    if conflicts:
        print(f"⚠️ 文件名冲突已追加序号：{len(set(conflicts))} 种")
    print(f"📊 报销导入 Excel：{reimb_path}（{len(reimb_rows)} 行）")
    print(f"📊 订单汇总 Excel：{order_path}（{len(order_rows)} 行）")
    print(f"统计：成功 {succ} / 失败 {fail}")


if __name__ == "__main__":
    main()
