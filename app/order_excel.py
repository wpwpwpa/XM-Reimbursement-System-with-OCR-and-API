"""淘宝订单 Excel 读取、按价格匹配、订单汇总 Excel 写出。

淘宝的时间与商品名来自订单 Excel（图片不识别），按图片识别出的价格匹配到 Excel 行。
表头自动识别常见写法（订单号/价格/下单时间/商品名称 等）。
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook, load_workbook

from order_fields import normalize_date

# 逻辑列 → 候选表头关键词（命中任一即采用该列）
COLUMN_ALIASES = {
    "price": ["实付金额", "价格", "金额", "实付款", "付款金额", "成交金额",
              "订单金额", "实付"],
    "order_time": ["下单时间", "下单日期", "付款时间", "成交时间", "订单时间",
                   "创建时间", "购买时间", "成交日期", "订单提交时间",
                   "提交时间", "订单创建时间", "拍下时间"],
    "product": ["商品名称", "宝贝标题", "商品", "标题", "商品标题", "商品信息"],
    "order_no": ["订单编号", "订单号", "订单id", "订单ID", "订单"],
}

PRICE_TOLERANCE = 5.0  # 价格匹配容差（元）；视觉模型金额识别常有 0.x~数元偏差


def _match_header(headers, keywords):
    """返回命中关键词最长的列索引（精确词如"实付金额"优先于泛词"金额"）；
    无命中返回 -1。"""
    best_idx = -1
    best_kw_len = 0
    for idx, h in enumerate(headers):
        hs = str(h).strip()
        for kw in keywords:
            if kw in hs and len(kw) > best_kw_len:
                best_idx = idx
                best_kw_len = len(kw)
    return best_idx


def _to_float(v):
    """把单元格值转 float；支持 数字 / "¥26.52" / "价格:26.52" 等写法。"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("¥", "").replace("￥", "").replace(",", "").strip()
    m = re.search(r"\d+(?:\.\d+)?", s)
    return float(m.group()) if m else None


def load_taobao_excel(path: str) -> dict:
    """读取淘宝订单 Excel。

    返回 {rows:[{price,order_time,product,order_no}], headers:[...]}。
    表头自动识别；识别不到的列在该行对应值为 None。
    read_only + data_only，兼容公式计算后的缓存值。
    """
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = list(next(rows_iter))
    except StopIteration:
        wb.close()
        return {"rows": [], "headers": []}
    idx = {key: _match_header(headers, kws)
           for key, kws in COLUMN_ALIASES.items()}

    out_rows = []
    for r in rows_iter:
        if r is None or all(c is None for c in r):
            continue
        row = {
            "price": _to_float(r[idx["price"]]) if idx["price"] >= 0 else None,
            "order_time": normalize_date(r[idx["order_time"]])
            if idx["order_time"] >= 0 else None,
            "product": str(r[idx["product"]]).strip()
            if (idx["product"] >= 0 and r[idx["product"]] is not None) else None,
            "order_no": str(r[idx["order_no"]]).strip()
            if (idx["order_no"] >= 0 and r[idx["order_no"]] is not None) else None,
        }
        out_rows.append(row)
    wb.close()
    return {"rows": out_rows, "headers": [str(h) for h in headers]}


def match_by_price(excel_data: dict, price: float | None):
    """按价格匹配淘宝 Excel 行。返回 (row_dict_or_None, ambiguous_bool)。

    容差内取价格最接近的一行；若有多行且差距都很小则标 ambiguous。
    """
    if price is None:
        return None, False
    cands = [(abs(r["price"] - price), r) for r in excel_data["rows"]
             if r["price"] is not None
             and abs(r["price"] - price) <= PRICE_TOLERANCE]
    if not cands:
        return None, False
    cands.sort(key=lambda x: x[0])
    # 次近与最近差距也极小 → 可能误配，标 ambiguous
    ambiguous = len(cands) > 1 and (cands[1][0] - cands[0][0]) < PRICE_TOLERANCE * 0.5
    return cands[0][1], ambiguous


def _lcs_ratio(a: str, b: str) -> float:
    """最长公共子串长度 / 较短串长度（0~1）。用于商品名相似度。"""
    shorter, longer = (a, b) if len(a) <= len(b) else (b, a)
    if not shorter:
        return 0.0
    best = 0
    for i in range(len(shorter)):
        for j in range(i + 1, len(shorter) + 1):
            if shorter[i:j] in longer:
                best = max(best, j - i)
    return best / len(shorter)


def match_by_product(excel_data: dict, product_name: str | None, ratio_thresh: float = 0.5):
    """按商品名称在 Excel 中匹配淘宝订单行。返回 (row_dict_or_None, ambiguous_bool)。

    视觉读出的商品名通常较短（如"一次性塑料滴管"），Excel 中较长
    （如"一次性塑料滴管实验室移液小吸管..."）。匹配策略：
    - 双向包含（视觉名⊆Excel名 或反之）→ 视为强匹配
    - 否则用最长公共子串相似度，>= ratio_thresh 才计入
    取相似度最高的一行；次高与最高极接近则标 ambiguous。
    """
    if not product_name or len(product_name.strip()) < 2:
        return None, False
    pn = product_name.strip()
    scored = []
    for r in excel_data["rows"]:
        en = (r.get("product") or "").strip()
        if not en:
            continue
        if pn in en or en in pn:
            score = 1.0
        else:
            score = _lcs_ratio(pn, en)
        if score >= ratio_thresh:
            scored.append((score, r))
    if not scored:
        return None, False
    scored.sort(key=lambda x: x[0], reverse=True)
    ambiguous = len(scored) > 1 and (scored[0][0] - scored[1][0]) < 0.15
    return scored[0][1], ambiguous


def build_order_rows(results, files) -> list[dict]:
    """由识别结果构造订单汇总行。results/files 须对齐。"""
    rows = []
    for f, parsed in zip(files, results):
        if not parsed:
            continue
        rows.append({
            "时间": parsed.get("order_time") or "",
            "商品名称": parsed.get("product_name") or "",
            "价格": parsed.get("amount") if parsed.get("amount") is not None else "",
            "平台": parsed.get("platform") or "",
            "原文件名": Path(f).name,
        })
    return rows


def write_order_excel(rows: list[dict], out_path: str) -> str:
    """写出订单汇总 Excel：时间 / 商品名称 / 价格 / 平台 / 原文件名。"""
    headers = ["时间", "商品名称", "价格", "平台", "原文件名"]
    wb = Workbook()
    ws = wb.active
    ws.title = "订单汇总"
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
    for col_idx, h in enumerate(headers, start=1):
        ws.column_dimensions[chr(64 + col_idx)].width = max(12, min(40, len(h) * 3))
    wb.save(out_path)
    return out_path
