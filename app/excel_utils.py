"""Excel 写出通用工具（被 order_excel / reimbursement_excel 共享）。"""
from __future__ import annotations

from openpyxl.utils import get_column_letter


def set_col_widths(ws, widths: list[float]) -> None:
    """按列序号设置列宽。

    用 openpyxl.utils.get_column_letter，支持 > 26 列；
    取代原先 chr(64 + c) 的写法（仅支持 A~Z，第 27 列会出错）。
    """
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
